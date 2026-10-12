"""Create schema-consistent copies of the quarterly source Excel workbooks."""

from __future__ import annotations

import re
from pathlib import Path

import pandas as pd
import xlrd
from openpyxl import load_workbook

from convert_pdfs import OUTPUT_COLUMNS, PROJECT_DIR, _write_filterable_workbook


SOURCE_DIR = PROJECT_DIR / "Junctoin Blacksite list (2004-2024)"
OUTPUT_DIR = PROJECT_DIR / "standardized_xlsx"
REGIONS = {
    "hong kong": "Hong Kong",
    "kowloon": "Kowloon",
    "new territories": "New Territories",
}
DISTRICT_REGIONS = {
    "central & west": "Hong Kong",
    "central & western": "Hong Kong",
    "central and western": "Hong Kong",
    "central and western district": "Hong Kong",
    "eastern": "Hong Kong",
    "southern": "Hong Kong",
    "wan chai": "Hong Kong",
    "kowloon city": "Kowloon",
    "kwun tong": "Kowloon",
    "sham shui po": "Kowloon",
    "wong tai sin": "Kowloon",
    "yau tsim mong": "Kowloon",
    "islands": "New Territories",
    "kwai tsing": "New Territories",
    "north": "New Territories",
    "sai kung": "New Territories",
    "sha tin": "New Territories",
    "shatin": "New Territories",
    "tai po": "New Territories",
    "tsuen wan": "New Territories",
    "tuen mun": "New Territories",
    "yuen long": "New Territories",
}
PERIOD_RE = re.compile(
    r"yearly data ending(?: at)? q([1-4])[- ]?(\d{2,4})", re.IGNORECASE
)
FILE_PERIOD_RE = re.compile(
    r"(?:Q([1-4])[-_ ]?(20\d{2}|\d{2})|Db_Q([1-4])_(\d{2}))",
    re.IGNORECASE,
)
SEVERITY_COLUMNS = ("fatal", "serious", "slight", "total")


def _text(value: object) -> str:
    return str(value).strip() if value is not None else ""


def _period_from_file(source_file: Path) -> tuple[int, int]:
    match = FILE_PERIOD_RE.search(source_file.name)
    if match is None:
        raise ValueError(f"Cannot determine year and quarter from {source_file.name}")
    quarter = int(match.group(1) or match.group(3))
    year = int(match.group(2) or match.group(4))
    if year < 100:
        year += 2000
    return year, quarter


def _period_year(value: str) -> int:
    year = int(value)
    return year + 2000 if year < 100 else year


def _section_labels(heading: str, region: str) -> tuple[str, str]:
    normalized = heading.casefold()
    if normalized in REGIONS:
        return REGIONS[normalized], ""
    return DISTRICT_REGIONS.get(normalized, region), heading.title()


def _record(
    source_file: Path,
    year: int,
    quarter: int,
    region: str,
    district: str,
    site_id: object,
    junction: object,
) -> dict[str, object]:
    street_1, street_2 = _split_junction(junction)
    return {
        "year": year,
        "quarter": quarter,
        "region": region or None,
        "district": district or None,
        "site_id": site_id,
        "street_1": street_1,
        "street_2": street_2 or None,
        "accidents_total": None,
        "accidents_pedestrians": None,
        "accidents_fatal_serious": None,
        "accidents_fatal": None,
        "accidents_serious": None,
        "accidents_slight": None,
        "casualties_fatal": None,
        "casualties_serious": None,
        "casualties_slight": None,
        "source_file": source_file.relative_to(PROJECT_DIR).as_posix(),
        "source_page": None,
    }


def _find_current_period_blocks(
    worksheet, header_row: int, year: int, quarter: int
) -> tuple[dict[str, dict[str, int]], int]:
    top_headers = [
        (cell.column, match)
        for cell in worksheet[header_row]
        if (match := PERIOD_RE.search(_text(cell.value))) is not None
    ]
    matches = [
        (column, match)
        for column, match in top_headers
        if int(match.group(1)) == quarter and int(match.group(2)) == year
    ]
    if not matches:
        return {}, header_row
    if len(matches) > 1:
        raise ValueError(
            f"Multiple current-period headers found in worksheet {worksheet.title!r}"
        )

    period_column, _ = matches[0]
    later_period_columns = [
        column for column, _ in top_headers if column > period_column
    ]
    block_end = min(later_period_columns, default=worksheet.max_column + 1)
    group_row = header_row + 1
    severity_row = header_row + 2
    groups: dict[str, dict[str, int]] = {}
    for cell in worksheet[group_row]:
        if not period_column <= cell.column < block_end:
            continue
        label = _text(cell.value).casefold()
        if label == "number of accidents":
            group_name = "accidents"
        elif label == "number of accidents involving pedestrians":
            group_name = "pedestrians"
        else:
            continue

        category_columns: dict[str, int] = {}
        for offset in range(len(SEVERITY_COLUMNS)):
            column = cell.column + offset
            category = _text(worksheet.cell(severity_row, column).value).casefold()
            if category in SEVERITY_COLUMNS:
                category_columns[category] = column
        if set(category_columns) != set(SEVERITY_COLUMNS):
            raise ValueError(
                f"Incomplete {group_name} metric headers in "
                f"{worksheet.title!r} at row {severity_row}"
            )
        groups[group_name] = category_columns

    return groups, severity_row


def _header_columns(worksheet) -> tuple[int, int, int, int] | None:
    for row in worksheet.iter_rows():
        values = {_text(cell.value).casefold(): cell.column for cell in row if cell.value}
        site_id_column = values.get("site id")
        junction_column = values.get("district & junction identification")
        if site_id_column is not None and junction_column is not None:
            return row[0].row, site_id_column, junction_column, junction_column + 1
    return None


def _is_five_year_only_table(worksheet, header_row: int) -> bool:
    period_header = any(
        "past five" in _text(cell.value).casefold()
        for cell in worksheet[header_row]
    )
    fatal_header = any(
        _text(cell.value).casefold() == "fatal"
        for row in worksheet.iter_rows(
            min_row=header_row + 1, max_row=header_row + 2
        )
        for cell in row
    )
    return period_header and fatal_header


def _metric_value(worksheet, row: int, column: int | None) -> int | None:
    if column is None:
        return None
    value = worksheet.cell(row, column).value
    if value is None or _text(value) == "":
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError(
            f"Expected a numeric metric at {worksheet.title}!{worksheet.cell(row, column).coordinate}, "
            f"found {value!r}"
        )
    if not float(value).is_integer():
        raise ValueError(
            f"Expected a whole-number metric at "
            f"{worksheet.title}!{worksheet.cell(row, column).coordinate}, found {value!r}"
        )
    return int(value)


def _xls_value(sheet, row: int, column: int) -> object:
    value = sheet.cell_value(row, column)
    return value if value != "" else None


def _xls_metric(sheet, row: int, column: int | None) -> int | None:
    if column is None:
        return None
    value = _xls_value(sheet, row, column)
    if value is None:
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError(
            f"Expected a numeric metric at {sheet.name}!R{row + 1}C{column + 1}, "
            f"found {value!r}"
        )
    if not float(value).is_integer():
        raise ValueError(
            f"Expected a whole-number metric at "
            f"{sheet.name}!R{row + 1}C{column + 1}, found {value!r}"
        )
    return int(value)


def _xls_row_texts(sheet, row: int) -> list[str]:
    return [_text(_xls_value(sheet, row, column)) for column in range(sheet.ncols)]


def _find_xls_headers(sheet) -> list[tuple[int, int | None, int]]:
    headers = []
    for row_number in range(sheet.nrows):
        values = [value.casefold() for value in _xls_row_texts(sheet, row_number)]
        site_id_column = next(
            (column for column, value in enumerate(values) if value == "site id"),
            None,
        )
        junction_column = next(
            (
                column
                for column, value in enumerate(values)
                if "junction identification" in value
            ),
            None,
        )
        if junction_column is not None:
            headers.append((row_number, site_id_column, junction_column))
    return headers


def _find_xls_metric_blocks(
    sheet, header_row: int, year: int, quarter: int
) -> tuple[dict[str, dict[str, int]], int] | None:
    period_headers = []
    for column, value in enumerate(_xls_row_texts(sheet, header_row)):
        match = PERIOD_RE.search(value)
        if match is not None:
            period_headers.append((column, match))
    current_headers = [
        (column, match)
        for column, match in period_headers
        if int(match.group(1)) == quarter and _period_year(match.group(2)) == year
    ]
    if period_headers and len(current_headers) != 1:
        return None
    if not period_headers and (year != 2009 or quarter != 4):
        return None

    first_metric_column = current_headers[0][0] if current_headers else 0
    later_period_columns = [
        column
        for column, _ in period_headers
        if column > first_metric_column
    ]
    block_end = min(later_period_columns, default=sheet.ncols)
    group_row = header_row + 1
    severity_row = header_row + 2
    if severity_row >= sheet.nrows:
        return None

    groups: dict[str, dict[str, int]] = {}
    for column, label in enumerate(_xls_row_texts(sheet, group_row)):
        if not first_metric_column <= column < block_end:
            continue
        normalized = label.casefold()
        if normalized == "number of accidents":
            group_name = "accidents"
        elif normalized == "number of accidents involving pedestrians":
            group_name = "pedestrians"
        else:
            continue

        metric_columns = {}
        for offset in range(len(SEVERITY_COLUMNS)):
            metric_column = column + offset
            if metric_column >= sheet.ncols:
                break
            severity = _text(_xls_value(sheet, severity_row, metric_column)).casefold()
            if severity in SEVERITY_COLUMNS:
                metric_columns[severity] = metric_column
        if set(metric_columns) == set(SEVERITY_COLUMNS):
            groups[group_name] = metric_columns
    if not groups:
        return None
    return groups, severity_row


def _is_legacy_data_row(
    sheet, row: int, metric_columns: set[int]
) -> bool:
    return any(
        isinstance(_xls_value(sheet, row, column), (int, float))
        and not isinstance(_xls_value(sheet, row, column), bool)
        for column in metric_columns
    )


def _extract_2009_legacy_table(source_file: Path, sheet, header_row: int) -> list[dict[str, object]]:
    year, quarter = _period_from_file(source_file)
    texts = [_xls_row_texts(sheet, row) for row in range(sheet.nrows)]
    subheaders = texts[header_row + 2]
    group_starts = []
    for column, label in enumerate(texts[header_row]):
        normalized = label.casefold()
        if normalized.startswith("no. of accidents"):
            group_starts.append((column, "accidents"))
        elif normalized.startswith("no. of accident"):
            group_starts.append((column, "accidents_by_severity"))
        elif normalized.startswith("no. of casualties"):
            group_starts.append((column, "casualties"))

    metric_positions: dict[str, int] = {}
    for column, label in enumerate(subheaders):
        normalized = label.casefold().replace(" ", "")
        group_name = next(
            (
                name
                for start, name in reversed(group_starts)
                if start <= column
            ),
            "",
        )
        if normalized == "total":
            if group_name == "accidents":
                metric_positions["accidents_total"] = column
        elif normalized == "ped":
            metric_positions["accidents_pedestrians"] = column
        elif normalized in {"fat/ser", "fat/ser."}:
            metric_positions["accidents_fatal_serious"] = column
        elif normalized in {"fatal", "serious", "slight"}:
            if group_name == "accidents_by_severity":
                metric_positions[f"accidents_{normalized}"] = column
            elif group_name == "casualties":
                metric_positions[f"casualties_{normalized}"] = column

    location_column = next(
        column
        for column, label in enumerate(texts[header_row])
        if label.casefold() == "junction identification"
    )
    site_id_column = next(
        column
        for column, label in enumerate(texts[header_row])
        if label.casefold() == "site id"
    )
    metric_columns = set(metric_positions.values())
    rows = []
    for row_number in range(header_row + 3, sheet.nrows):
        junction = _xls_value(sheet, row_number, location_column)
        if junction is None or not _is_legacy_data_row(sheet, row_number, metric_columns):
            continue
        record = _record(
            source_file,
            year,
            quarter,
            "",
            "",
            _xls_value(sheet, row_number, site_id_column),
            junction,
        )
        for name, column in metric_positions.items():
            record[name] = _xls_metric(sheet, row_number, column)
        rows.append(record)
    return rows


def _extract_xls_rows(source_file: Path) -> list[dict[str, object]]:
    year, quarter = _period_from_file(source_file)
    workbook = xlrd.open_workbook(source_file, on_demand=True)
    rows: list[dict[str, object]] = []
    try:
        for sheet in workbook.sheets():
            header_rows = _find_xls_headers(sheet)
            special_2009 = next(
                (
                    header_row
                    for header_row, site_id_column, _ in header_rows
                    if year == 2009
                    and quarter < 4
                    and site_id_column is not None
                    and any(
                        value.casefold() == "rank"
                        for value in _xls_row_texts(sheet, header_row)
                    )
                ),
                None,
            )
            if special_2009 is not None:
                rows.extend(
                    _extract_2009_legacy_table(
                        source_file, sheet, special_2009
                    )
                )
                continue

            for header_row, site_id_column, junction_header_column in header_rows:
                junction_column = (
                    junction_header_column + 1
                    if site_id_column is not None
                    else junction_header_column
                )
                if junction_column >= sheet.ncols:
                    continue
                metric_block = _find_xls_metric_blocks(
                    sheet, header_row, year, quarter
                )
                if metric_block is None:
                    continue
                groups, last_header_row = metric_block
                metric_columns = {
                    column
                    for group in groups.values()
                    for column in group.values()
                }
                district = ""
                region = ""
                for row_number in range(last_header_row + 1, sheet.nrows):
                    site_id = (
                        _xls_value(sheet, row_number, site_id_column)
                        if site_id_column is not None
                        else None
                    )
                    heading_or_code = _xls_value(
                        sheet, row_number, junction_header_column
                    )
                    junction = _xls_value(sheet, row_number, junction_column)
                    if not _is_legacy_data_row(sheet, row_number, metric_columns):
                        if (
                            heading_or_code is not None
                            and (
                                junction is None
                                or _text(junction) == ""
                                or (
                                    site_id_column is None
                                    and junction_header_column == junction_column
                                    and not isinstance(heading_or_code, (int, float))
                                )
                            )
                        ):
                            region, district = _section_labels(
                                _text(heading_or_code), region
                            )
                        continue
                    if junction is None or _text(junction) == "":
                        continue
                    record = _record(
                        source_file,
                        year,
                        quarter,
                        region,
                        district,
                        site_id,
                        junction,
                    )
                    accident_metrics = groups.get("accidents")
                    if accident_metrics is not None:
                        for severity in SEVERITY_COLUMNS:
                            value = _xls_metric(
                                sheet, row_number, accident_metrics[severity]
                            )
                            if severity == "total":
                                record["accidents_total"] = value
                            else:
                                record[f"accidents_{severity}"] = value
                        fatal = record["accidents_fatal"]
                        serious = record["accidents_serious"]
                        if fatal is not None and serious is not None:
                            record["accidents_fatal_serious"] = fatal + serious
                    pedestrian_metrics = groups.get("pedestrians")
                    if pedestrian_metrics is not None:
                        record["accidents_pedestrians"] = _xls_metric(
                            sheet, row_number, pedestrian_metrics["total"]
                        )
                    rows.append(record)
                break

            for header_row, site_id_column, junction_header_column in header_rows:
                near_header = range(max(0, header_row - 1), min(sheet.nrows, header_row + 2))
                has_five_year_heading = any(
                    "past five" in value.casefold()
                    for row_number in near_header
                    for value in _xls_row_texts(sheet, row_number)
                )
                if not has_five_year_heading:
                    continue
                fatal_row = next(
                    (
                        row_number
                        for row_number in range(header_row + 1, min(sheet.nrows, header_row + 5))
                        if any(
                            value.casefold() == "fatal"
                            for value in _xls_row_texts(sheet, row_number)
                        )
                    ),
                    None,
                )
                if fatal_row is None:
                    continue
                junction_column = (
                    junction_header_column + 1
                    if site_id_column is not None
                    else junction_header_column
                )
                district = ""
                region = ""
                for row_number in range(fatal_row + 1, sheet.nrows):
                    site_id = (
                        _xls_value(sheet, row_number, site_id_column)
                        if site_id_column is not None
                        else None
                    )
                    heading_or_code = _xls_value(
                        sheet, row_number, junction_header_column
                    )
                    junction = _xls_value(sheet, row_number, junction_column)
                    metric_cells = [
                        _xls_value(sheet, row_number, column)
                        for column in range(junction_column + 1, sheet.ncols)
                    ]
                    has_number = any(
                        isinstance(value, (int, float))
                        and not isinstance(value, bool)
                        for value in metric_cells
                    )
                    if not has_number:
                        if heading_or_code is not None and (
                            junction is None or _text(junction) == ""
                            or (
                                site_id_column is None
                                and junction_header_column == junction_column
                                and not isinstance(heading_or_code, (int, float))
                            )
                        ):
                            region, district = _section_labels(
                                _text(heading_or_code), region
                            )
                        continue
                    if junction is None or _text(junction) == "":
                        continue
                    rows.append(
                        _record(
                            source_file,
                            year,
                            quarter,
                            region,
                            district,
                            site_id,
                            junction,
                        )
                    )
    finally:
        workbook.release_resources()
    if not rows:
        raise ValueError(
            f"No current-period or five-year junction records found in {source_file.name}"
        )
    return rows


def _split_junction(value: object) -> tuple[str, str]:
    parts = _text(value).split(",", maxsplit=1)
    return parts[0].strip(), parts[1].strip() if len(parts) == 2 else ""


def _extract_rows(source_file: Path) -> list[dict[str, object]]:
    year, quarter = _period_from_file(source_file)
    rows: list[dict[str, object]] = []
    workbook = load_workbook(source_file, read_only=False, data_only=True)
    try:
        for worksheet in workbook.worksheets:
            header = _header_columns(worksheet)
            if header is None:
                continue
            header_row, site_id_column, district_column, junction_column = header
            groups, last_header_row = _find_current_period_blocks(
                worksheet, header_row, year, quarter
            )
            five_year_only = not groups and _is_five_year_only_table(
                worksheet, header_row
            )
            if not groups and not five_year_only:
                continue

            region = ""
            district = ""
            first_data_row = (
                last_header_row + 1 if groups else header_row + 1
            )
            for row_number in range(first_data_row, worksheet.max_row + 1):
                site_id = worksheet.cell(row_number, site_id_column).value
                district_or_group = worksheet.cell(row_number, district_column).value
                junction = worksheet.cell(row_number, junction_column).value

                if site_id is None and junction is None and district_or_group is not None:
                    heading = _text(district_or_group)
                    region, district = _section_labels(heading, region)
                    continue

                if junction is None or _text(junction) == "":
                    continue

                record = _record(
                    source_file,
                    year,
                    quarter,
                    region,
                    district,
                    site_id,
                    junction,
                )

                accident_metrics = groups.get("accidents")
                if accident_metrics is not None:
                    for severity in SEVERITY_COLUMNS:
                        value = _metric_value(
                            worksheet, row_number, accident_metrics[severity]
                        )
                        if severity == "total":
                            record["accidents_total"] = value
                        else:
                            record[f"accidents_{severity}"] = value
                    fatal = record["accidents_fatal"]
                    serious = record["accidents_serious"]
                    if fatal is not None and serious is not None:
                        record["accidents_fatal_serious"] = fatal + serious

                pedestrian_metrics = groups.get("pedestrians")
                if pedestrian_metrics is not None:
                    record["accidents_pedestrians"] = _metric_value(
                        worksheet, row_number, pedestrian_metrics["total"]
                    )

                rows.append(record)
    finally:
        workbook.close()

    if not rows:
        raise ValueError(
            f"No current-quarter annual records were found in {source_file.name}"
        )
    return rows


def standardize_all_excels() -> list[Path]:
    source_files = sorted(
        path
        for year_directory in SOURCE_DIR.iterdir()
        if year_directory.is_dir() and year_directory.name.isdigit()
        and 2009 <= int(year_directory.name) <= 2024
        for path in year_directory.iterdir()
        if path.is_file() and path.suffix.casefold() in {".xls", ".xlsx"}
    )
    if not source_files:
        raise FileNotFoundError(f"No quarterly Excel workbooks found under {SOURCE_DIR}")

    standardized_files = []
    for source_file in source_files:
        records = (
            _extract_xls_rows(source_file)
            if source_file.suffix.casefold() == ".xls"
            else _extract_rows(source_file)
        )
        data = pd.DataFrame(records, columns=OUTPUT_COLUMNS)
        year, quarter = _period_from_file(source_file)
        destination = (
            OUTPUT_DIR
            / source_file.parent.name
            / f"{source_file.stem}.xlsx"
        )
        table_name = f"Standardized_{year}_Q{quarter}"
        _write_filterable_workbook(data, destination, table_name)
        standardized_files.append(destination)
        print(f"Standardized {source_file.name}: {len(data)} annual records")

    return standardized_files


if __name__ == "__main__":
    files = standardize_all_excels()
    print(f"\nCreated {len(files)} standardized copies under {OUTPUT_DIR}")

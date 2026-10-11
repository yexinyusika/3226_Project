"""Convert the 2004-2008 Junction Blacksite PDFs into filterable Excel tables."""

from __future__ import annotations

import re
import shutil
from dataclasses import dataclass
from pathlib import Path

import pandas as pd
import pdfplumber
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo


PROJECT_DIR = Path(__file__).resolve().parent
SOURCE_DIR = PROJECT_DIR / "Junctoin Blacksite list (2004-2024)"
OUTPUT_DIR = PROJECT_DIR / "converted_pdf_xlsx"

METRIC_COLUMNS = [
    "accidents_total",
    "accidents_pedestrians",
    "accidents_fatal_serious",
    "accidents_fatal",
    "accidents_serious",
    "accidents_slight",
    "casualties_fatal",
    "casualties_serious",
    "casualties_slight",
]
OUTPUT_COLUMNS = [
    "year",
    "quarter",
    "region",
    "district",
    "site_id",
    "street_1",
    "street_2",
    *METRIC_COLUMNS,
    "source_file",
    "source_page",
]

REFERENCE_PAGE_WIDTH = 842


@dataclass(frozen=True)
class Word:
    text: str
    x0: float
    x1: float
    top: float


def _group_words(words: list[Word], tolerance: float = 3.0) -> list[list[Word]]:
    lines: list[list[Word]] = []
    for word in sorted(words, key=lambda item: (item.top, item.x0)):
        if not lines or abs(word.top - lines[-1][0].top) > tolerance:
            lines.append([word])
        else:
            lines[-1].append(word)
    return [sorted(line, key=lambda item: item.x0) for line in lines]


def _ocr_words(page: pdfplumber.page.Page, source_file: Path) -> list[Word]:
    if shutil.which("tesseract") is None:
        raise RuntimeError(
            f"{source_file.name} uses an unreadable embedded font. Install Tesseract "
            "OCR (for example, `sudo apt-get install tesseract-ocr`) and the "
            "`pytesseract` Python dependency, then rerun the conversion."
        )

    import pytesseract
    from pytesseract import Output

    image = page.to_image(resolution=300).original
    scale = image.width / page.width
    data = pytesseract.image_to_data(
        image,
        config="--psm 6",
        output_type=Output.DICT,
    )
    words = []
    for index, text in enumerate(data["text"]):
        text = text.strip()
        if not text or float(data["conf"][index]) < 20:
            continue
        left = float(data["left"][index])
        width = float(data["width"][index])
        top = float(data["top"][index])
        words.append(
            Word(
                text=text,
                x0=left / scale,
                x1=(left + width) / scale,
                top=top / scale,
            )
        )
    return words


def _page_words(
    page: pdfplumber.page.Page, source_file: Path
) -> tuple[list[Word], bool]:
    text = page.extract_text() or ""
    readable_header = re.search(
        r"district\s+board.*?1st\s+street.*?2nd\s+street",
        text,
        flags=re.IGNORECASE | re.DOTALL,
    )
    if readable_header:
        words = [
            Word(
                text=word["text"],
                x0=float(word["x0"]),
                x1=float(word["x1"]),
                top=float(word["top"]),
            )
            for word in page.extract_words(x_tolerance=2, y_tolerance=3)
        ]
        return words, False
    return _ocr_words(page, source_file), True


def _region_from_page(words: list[Word]) -> str | None:
    text = " ".join(word.text for word in sorted(words, key=lambda item: (item.top, item.x0)))
    match = re.search(
        r"(?:junction\s*[-–]?\s*blacksite\s+list\s+for)\s+(.+?)\s+region\b",
        text,
        flags=re.IGNORECASE,
    )
    return match.group(1).strip().title() if match else None


def _column_headers(words: list[Word]) -> tuple[float, float, float]:
    header_words = [
        word
        for word in words
        if 100 <= word.top <= 150
        and (
            word.text.lower() in {"1st", "ist", "ast", "lst", "2nd"}
            or word.text.lower().startswith("total")
        )
    ]
    first_street = next(
        (
            word
            for word in header_words
            if word.text.lower() in {"1st", "ist", "ast", "lst"}
        ),
        None,
    )
    second_street = next(
        (word for word in header_words if word.text.lower() == "2nd"), None
    )
    total = next(
        (word for word in header_words if word.text.lower() == "total"), None
    )
    if first_street is None or second_street is None or total is None:
        raise ValueError("Could not locate the table column headings in the PDF page")
    return first_street.x0, second_street.x0, total.x0


def _numeric_value(text: str) -> int | None:
    if not re.fullmatch(r"[0-9Il|]+", text):
        return None
    normalized = text.translate(str.maketrans({"I": "1", "l": "1", "|": "1"}))
    return int(normalized)


def _parse_pdf_page(
    page: pdfplumber.page.Page,
    page_number: int,
    source_file: Path,
    year: int,
    quarter: int,
) -> list[dict[str, object]]:
    page_text = page.extract_text() or ""
    has_table_header = re.search(
        r"district\s+board.*?1st\s+street.*?2nd\s+street",
        page_text,
        flags=re.IGNORECASE | re.DOTALL,
    )
    has_report_summary = re.search(
        r"total\s+number\s+of\s+blacksites\s+in\s+all\s+region",
        page_text,
        flags=re.IGNORECASE,
    )
    if has_report_summary and not has_table_header:
        return []
    words, used_ocr = _page_words(page, source_file)
    lines = _group_words(words, tolerance=4.0 if used_ocr else 3.0)
    normalized_ocr_text = ""
    if used_ocr and not has_table_header:
        ocr_text = " ".join(
            word.text
            for word in sorted(words, key=lambda item: (item.top, item.x0))
        )
        normalized_ocr_text = re.sub(r"[^a-z]+", " ", ocr_text.lower())
    try:
        street_1_start, street_2_start, numeric_start = _column_headers(words)
    except ValueError:
        if not used_ocr:
            raise
        has_ocr_footer = bool(
            re.search(
                r"total number of blacks?ites? in all region|end of report",
                normalized_ocr_text,
            )
        )
        has_table_rows = any(
            sum(
                word.x0 >= 0.53 * page.width
                and _numeric_value(word.text) is not None
                for word in line
            )
            >= len(METRIC_COLUMNS)
            for line in lines
        )
        if has_ocr_footer and not has_table_rows:
            return []
        if not has_table_rows:
            raise ValueError(
                f"Could not locate table headings or records in "
                f"{source_file.name}, page {page_number}"
            )
        street_1_start = 0.128 * page.width
        street_2_start = 0.325 * page.width
        numeric_start = 0.536 * page.width

    region = _region_from_page(words)
    current_district: str | None = None
    page_rows: list[dict[str, object]] = []
    width_scale = page.width / REFERENCE_PAGE_WIDTH
    street_1_start = min(street_1_start, 0.12 * page.width)
    numeric_start -= 2
    street_2_end = numeric_start

    for line in lines:
        if not line or line[0].top < 130 * width_scale:
            continue
        if line[0].top > page.height - 10 * width_scale:
            continue

        line_text = " ".join(word.text for word in line).strip()
        if line[0].x0 < 0.12 * page.width and not any(
            re.fullmatch(r"\d+", word.text) for word in line
        ):
            district_candidate = line_text.strip()
            if (
                re.fullmatch(r"[A-Za-z][A-Za-z &'()./-]*", district_candidate)
                and not district_candidate.lower().startswith(("number of", "blacksites"))
            ):
                current_district = district_candidate.title()
                continue

        street_1_words = [
            word for word in line if street_1_start <= word.x0 < street_2_start
        ]
        street_2_words = [
            word for word in line if street_2_start <= word.x0 < street_2_end
        ]
        if not street_1_words or not street_2_words:
            continue
        numeric_words = [
            word
            for word in line
            if word.x0 >= numeric_start and _numeric_value(word.text) is not None
        ]
        if not numeric_words:
            continue

        numeric_words.sort(key=lambda word: word.x0)
        if len(numeric_words) != len(METRIC_COLUMNS):
            raise ValueError(
                f"Expected all {len(METRIC_COLUMNS)} statistics in "
                f"{source_file.name}, page {page_number}, but found "
                f"{len(numeric_words)}: {line_text}"
            )
        values = [_numeric_value(word.text) for word in numeric_words]

        row: dict[str, object] = {
            "year": year,
            "quarter": quarter,
            "region": region,
            "district": current_district,
            "site_id": None,
            "street_1": " ".join(word.text for word in street_1_words).strip(),
            "street_2": " ".join(word.text for word in street_2_words).strip(),
            "source_file": str(source_file.relative_to(PROJECT_DIR)),
            "source_page": page_number,
        }
        row.update(dict(zip(METRIC_COLUMNS, values, strict=True)))
        page_rows.append(row)

    return page_rows


def _parse_pdf(source_file: Path) -> pd.DataFrame:
    source_file = source_file.resolve()
    quarter_match = re.search(r"Q([1-4])", source_file.stem, flags=re.IGNORECASE)
    if quarter_match is None:
        raise ValueError(f"Cannot determine the quarter from {source_file.name}")
    try:
        year = int(source_file.parent.name)
    except ValueError as error:
        raise ValueError(f"Cannot determine the year from {source_file}") from error
    quarter = int(quarter_match.group(1))

    rows = []
    with pdfplumber.open(source_file) as pdf:
        for page_number, page in enumerate(pdf.pages, start=1):
            rows.extend(
                _parse_pdf_page(page, page_number, source_file, year, quarter)
            )
        document_text = "\n".join(page.extract_text() or "" for page in pdf.pages)
        total_match = re.search(
            r"total\s+number\s+of\s+blacksites\s+in\s+all\s+region\s*:\s*(\d+)",
            document_text,
            flags=re.IGNORECASE,
        )
        if total_match is None:
            summary_words = _ocr_words(pdf.pages[-1], source_file)
            summary_text = " ".join(
                word.text
                for word in sorted(summary_words, key=lambda item: (item.top, item.x0))
            )
            total_match = re.search(
                r"total\s+number\s+of\s+blacksites\s+in\s+all\s+region\D{0,20}(\d+)",
                summary_text,
                flags=re.IGNORECASE,
            )
        if total_match is None:
            raise ValueError(
                f"Could not read the printed all-regions record total in "
                f"{source_file.name}; check the PDF or OCR quality."
            )
        expected_total = int(total_match.group(1))
        if len(rows) != expected_total:
            raise ValueError(
                f"Extracted {len(rows)} records from {source_file.name}, but the "
                f"PDF reports {expected_total}; no workbooks were written."
            )
    if not rows:
        raise ValueError(f"No junction records were extracted from {source_file}")
    data = pd.DataFrame(rows, columns=OUTPUT_COLUMNS)
    data["district"] = data.groupby("region", sort=False)["district"].ffill()
    if data["region"].isna().any() or data["district"].isna().any():
        raise ValueError(
            f"Some records in {source_file.name} are missing their region or "
            "district after carrying forward headings across page breaks."
        )
    return data


def _write_filterable_workbook(
    data: pd.DataFrame, destination: Path, table_name: str
) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    data.to_excel(destination, sheet_name="data", index=False, na_rep="")

    workbook = load_workbook(destination)
    worksheet = workbook["data"]
    worksheet.freeze_panes = "A2"
    worksheet.sheet_view.showGridLines = False
    if worksheet.max_row > 1:
        table = Table(displayName=table_name, ref=worksheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        worksheet.add_table(table)
    worksheet.auto_filter.ref = worksheet.dimensions
    for cells in worksheet.columns:
        header = str(cells[0].value or "")
        worksheet.column_dimensions[cells[0].column_letter].width = min(
            max(max((len(str(cell.value or "")) for cell in cells), default=0) + 2, 12),
            42,
        )
        if header in {"street_1", "street_2", "source_file"}:
            worksheet.column_dimensions[cells[0].column_letter].width = 32
    workbook.save(destination)


def convert_all_pdfs() -> tuple[list[Path], Path]:
    pdf_files = sorted(SOURCE_DIR.glob("200[4-8]/*.pdf"))
    if not pdf_files:
        raise FileNotFoundError(f"No 2004-2008 PDF files found under {SOURCE_DIR}")

    quarterly_data: list[tuple[Path, pd.DataFrame]] = []
    for source_file in pdf_files:
        data = _parse_pdf(source_file)
        quarterly_data.append((source_file, data))
        print(f"Parsed {source_file.name}: {len(data)} junction records")

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    quarterly_workbooks = []
    master_parts = []
    for source_file, data in quarterly_data:
        destination = (
            OUTPUT_DIR
            / source_file.parent.name
            / f"{source_file.stem}.xlsx"
        )
        table_name = f"Blacksite_{source_file.parent.name}_Q{int(data['quarter'].iloc[0])}"
        _write_filterable_workbook(data, destination, table_name)
        quarterly_workbooks.append(destination)
        master_parts.append(data)

    master = pd.concat(master_parts, ignore_index=True)
    master_destination = OUTPUT_DIR / "Junction_Blacksite_Master_2004-2008.xlsx"
    _write_filterable_workbook(master, master_destination, "Blacksite_Master")
    return quarterly_workbooks, master_destination


if __name__ == "__main__":
    quarterly_files, master_file = convert_all_pdfs()
    print(
        f"\nCreated {len(quarterly_files)} quarterly workbooks and "
        f"one master workbook in {OUTPUT_DIR}"
    )
    print(f"Master workbook: {master_file}")
